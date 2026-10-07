from copy import deepcopy
from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"D:/Schooling/PT&TK HTTT/ASSIGNMENT/A6_V1_03_NguyenTienDat017.docx")
OUTPUT_DIR = ROOT / "output/docx"
OUTPUT = OUTPUT_DIR / SOURCE.name
ASSETS = ROOT / "tmp/docx-assets"


def set_font(run, name="Times New Roman", size=10, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def replace_paragraph(doc, startswith, text):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(startswith):
            paragraph.clear()
            run = paragraph.add_run(text)
            set_font(run)
            return paragraph
    raise ValueError(f"Paragraph not found: {startswith}")


def set_cell_text(cell, text, bold=False, color="000000", size=8.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def format_table(table, widths=None, header_fill="1F4E78"):
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[col_index])
            if row_index == 0:
                shade_cell(cell, header_fill)
                for run in cell.paragraphs[0].runs:
                    set_font(run, size=8.5, bold=True, color="FFFFFF")
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif row_index % 2 == 0:
                shade_cell(cell, "F3F6F9")
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], str(value))
    format_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.08
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.15)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run("• ")
    paragraph.add_run(text)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = False
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, size=8.5, italic=True)
    return paragraph


def add_picture(doc, path, width):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    return paragraph


def update_existing_content(doc):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name in {"Heading 1", "Heading 2"} and text.endswith(";"):
            paragraph.text = text[:-1]

    replace_paragraph(doc, "The prototype uses keyword matching", (
        "The final prototype combines deterministic keyword matching with real browser microphone input and pretrained visual embeddings. "
        "Text and recognized speech share bilingual Vietnamese-English normalization. Image search uses CLIP ViT-B/32 to encode uploaded images and catalog photographs, then ranks products by cosine similarity. "
        "The implementation therefore covers the required simulated voice workflow in the CLI and extends it with microphone recognition and arbitrary image upload in the web interface."
    ))
    replace_paragraph(doc, "This report presents the analysis", (
        "This report presents the analysis, design, and implementation of a prototype multimodal search system for e-commerce. "
        "The system supports product searches through text, browser microphone input, simulated voice transcripts for repeatable CLI demonstrations, and uploaded product images. "
        "These inputs follow a unified process: query representation, candidate retrieval, ranking, and result presentation."
    ))
    replace_paragraph(doc, "Voice search: A customer", (
        "Voice search: On the web, the customer records a request through the microphone and the browser converts speech to a transcript. "
        "The transcript remains editable and then follows the same normalization and retrieval process as text search. The CLI retains the assignment's simulated transcript workflow for repeatable demonstration."
    ))
    replace_paragraph(doc, "Image search: A customer", (
        "Image search: A customer uploads a JPEG, PNG, or WebP image. A pretrained CLIP ViT-B/32 vision encoder converts its pixels into a 512-dimensional semantic embedding. "
        "The system compares this vector with embeddings inferred from the catalog photographs and returns the nearest products."
    ))
    replace_paragraph(doc, "The product dataset contains at least ten", (
        "The product repository contains 40 products across four balanced categories: shoes, bags and backpacks, clothing, and drinkware. "
        "Each category contains ten products. Every record has a stable identifier, English and Vietnamese names, category, color, price, stock quantity, description, and a unique credited photograph. "
        "Search results display product details, rank, and score. Empty input, invalid uploads, unsupported formats, and no-match queries receive clear feedback."
    ))
    replace_paragraph(doc, "The dataset shall contain at least ten", (
        "The dataset contains 40 products with consistent identifiers, balanced category counts, unique image paths, and credited source URLs. "
        "Catalog and query embeddings have 512 dimensions and are produced by the same CLIP preprocessing pipeline. "
        "The interface identifies browser speech recognition as an external service and reports the image model used for retrieval."
    ))
    replace_paragraph(doc, "Customer is the primary actor", (
        "Customer is the primary actor. Customers search by text, microphone, or uploaded image and inspect ranked results. "
        "The CLI also supports the required simulated voice transcript. In the extended design scope, customers can search for and view orders."
    ))
    replace_paragraph(doc, "Search by Keyword, Search by Voice", (
        "Search by Keyword, Search by Voice, and Search by Image specialize Search Product through generalization relationships. "
        "The web realization accepts real microphone input and arbitrary image uploads, while the CLI preserves deterministic text, simulated voice, and image demonstrations."
    ))
    replace_paragraph(doc, "The system follows a three-layer architecture", (
        "The system follows a three-layer architecture that separates user interaction, query and ranking logic, and data access. "
        "This separation allowed the initial artificial image vectors to be replaced by CLIP without changing the common query schema or result-ranking interface."
    ))
    replace_paragraph(doc, "SearchUI uses the input components", (
        "SearchUI and the FastAPI presentation adapter collect input, invoke application services, and render results. QueryService normalizes text and constructs the common query object. "
        "SearchService retrieves candidates from ProductRepository and uses VectorIndex plus RankingService for scored output. "
        "For image requests, ImageService calls CLIPImageEncoder; catalog images are indexed at startup and uploaded images are processed in memory."
    ))
    replace_paragraph(doc, "The architecture represents logical responsibilities", (
        "The diagrams represent logical responsibilities. In the final Python implementation, the image-storage responsibility is realized by local image files, while CLIPImageEncoder performs feature extraction. "
        "Order components remain an extended design scope and are not part of the product-search prototype."
    ))
    replace_paragraph(doc, "The sequence diagram describes how", (
        "The sequence diagram documents the assignment's required simulated voice scenario. SearchUI receives a supplied transcript, validates it through SpeechService, constructs a normalized query, retrieves candidates, ranks them, and displays the result. "
        "The web interface adds microphone recognition before this sequence and sends only the resulting transcript to the Python API."
    ))

    # Functional requirements table.
    replacements = {
        "FR02": "The system shall accept a supplied transcript for repeatable voice-search demonstration and a browser-recognized microphone transcript in the web interface.",
        "FR03": "The system shall accept an arbitrary JPEG, PNG, or WebP image and perform similarity search using pretrained CLIP embeddings.",
        "FR05": "The system shall validate inputs, uploaded file type and size, and return clear feedback for invalid requests or no matches.",
        "FR06": "The system shall retrieve candidate products and rank them by a deterministic relevance or cosine-similarity score.",
    }
    for row in doc.tables[1].rows[1:]:
        key = row.cells[0].text.strip()
        if key in replacements:
            set_cell_text(row.cells[1], replacements[key])

    nfr_replacements = {
        "NFR02": "Speech recognition and image encoders shall remain replaceable behind stable service interfaces.",
        "NFR04": "Identical inputs and model files shall produce deterministic ranking, including a defined product-ID tie rule.",
        "NFR05": "The interface shall identify browser speech processing, local image inference, accepted upload formats, and prototype limitations.",
        "NFR06": "Uploaded images shall be processed in memory, limited to 8 MB, and never stored by the server.",
    }
    for row in doc.tables[2].rows[1:]:
        key = row.cells[0].text.strip().split()[0]
        if key in nfr_replacements:
            set_cell_text(row.cells[1], nfr_replacements[key])

    io_rows = doc.tables[3].rows
    set_cell_text(io_rows[2].cells[1], "Microphone audio or supplied transcript")
    set_cell_text(io_rows[2].cells[2], "Browser speech recognition, transcript validation, then bilingual text search")
    set_cell_text(io_rows[3].cells[1], "JPEG, PNG, WebP upload or demo image")
    set_cell_text(io_rows[3].cells[2], "CLIP ViT-B/32 inference and cosine similarity")

    use_case_rows = doc.tables[4].rows
    set_cell_text(use_case_rows[3].cells[1], "Find products using a microphone transcript or a supplied transcript for the CLI demonstration.")
    set_cell_text(use_case_rows[4].cells[1], "Find visually similar products from an arbitrary supported image upload using CLIP embeddings.")

    scenario_rows = doc.tables[5].rows
    set_cell_text(scenario_rows[4].cells[1], "1. The customer records speech with the web microphone or supplies a CLI transcript. 2. The browser or CLI produces a transcript. 3. The system constructs a query. 4. Candidate products are retrieved and ranked. 5. Results and scores are displayed.")
    set_cell_text(scenario_rows[7].cells[1], "The web uses browser microphone recognition; the CLI uses a supplied transcript so the required simulated demonstration is repeatable.")

    presentation_rows = doc.tables[6].rows
    set_cell_text(presentation_rows[2].cells[1], "Capture microphone input in the browser or accept a supplied CLI transcript.")
    set_cell_text(presentation_rows[3].cells[1], "Accept a JPEG, PNG, or WebP upload of at most 8 MB, or select a registered demo image.")

    app_rows = doc.tables[7].rows
    set_cell_text(app_rows[2].cells[1], "Validate the transcript supplied by the browser or CLI before query normalization.")
    set_cell_text(app_rows[3].cells[1], "Use CLIPImageEncoder to infer a normalized 512-dimensional vector from image pixels.")
    set_cell_text(app_rows[4].cells[1], "Retrieve text candidates or compare image vectors against the complete catalog index.")

    data_rows = doc.tables[8].rows
    set_cell_text(data_rows[3].cells[1], "Store normalized CLIP vectors for all 40 catalog products during the application process.")
    set_cell_text(data_rows[5].cells[1], "Store order records in the extended design scope; this component is not implemented in the search prototype.")
    set_cell_text(data_rows[6].cells[1], "Store product images as unique local PNG files with source and license metadata; the CLIP ONNX model is loaded as a separate runtime artifact.")

    participant_rows = doc.tables[9].rows
    set_cell_text(participant_rows[3].cells[1], "Validate and return the supplied transcript in the deterministic CLI sequence.")


def append_sections(doc):
    doc.add_page_break()
    add_heading(doc, "7. Python Implementation", 1)
    add_body(doc, (
        "The implementation uses Python 3.13 and preserves the three-layer design. The CLI composition root and the FastAPI web adapter build the same application services, so both interfaces execute the same repository, query, retrieval, and ranking logic. "
        "The web adds HTML, CSS, and JavaScript for microphone control, image upload, validation messages, and bilingual presentation."
    ))

    add_heading(doc, "7.1. Project Structure", 2)
    add_table(doc, ["Package or file", "Implementation responsibility"], [
        ("data", "ProductRepository, VectorIndex, and the ONNX image encoder."),
        ("application", "Query normalization, speech transcript validation, image encoding, candidate retrieval, cosine similarity, and ranking."),
        ("presentation", "CLI input coordination and formatted demonstration output."),
        ("web", "FastAPI endpoints and the browser interface for text, microphone, image upload, catalog, and diagnostics."),
        ("datasets", "Forty product records, unique product photographs, demo query images, evaluation queries, and photo-source metadata."),
        ("models", "CLIP ViT-B/32 ONNX INT8 vision encoder and the retained MobileNetV2 comparison model."),
        ("tests and scripts", "Unit, API, browser, catalog-photo, retrieval, and reproducible evaluation checks."),
    ], widths=[1.55, 5.05])

    add_heading(doc, "7.2. Product Repository and Dataset", 2)
    add_body(doc, (
        "ProductRepository loads and validates datasets/products.json. The catalog contains 40 products, with ten records in each of four categories. "
        "Every product has English and Vietnamese names and a distinct local photograph. photo_sources.json records a separate source URL for each image. "
        "The repository returns copies of records so presentation code cannot mutate stored data."
    ))

    add_heading(doc, "7.3. Search and Ranking Services", 2)
    add_body(doc, (
        "QueryService creates a common query object with type, normalized text or embedding, and filters. SearchService selects the retrieval path by query type. "
        "RankingService validates top_k, sorts by descending score, uses ascending product ID for deterministic ties, and returns only the requested number of products."
    ))
    add_body(doc, (
        "The FastAPI adapter exposes catalog, product-detail, text-search, voice-search, and image-search endpoints. Image uploads are restricted to JPEG, PNG, and WebP, limited to 8 MB, decoded before inference, and processed without permanent storage. "
        "The API hides internal paths and returns controlled 4xx or 5xx responses instead of private tracebacks."
    ))

    add_heading(doc, "7.4. Running the Program", 2)
    add_table(doc, ["Purpose", "Command"], [
        ("Install dependencies", r".\.venv-web\Scripts\python.exe -m pip install -r requirements-web.txt"),
        ("Run CLI demonstration", r".\.venv-web\Scripts\python.exe main.py --demo"),
        ("Run web application", r".\.venv-web\Scripts\python.exe -m uvicorn web.app:app --host 127.0.0.1 --port 8011"),
        ("Run automated tests", r".\.venv-web\Scripts\python.exe -m unittest discover -s tests -v"),
        ("Compare image models", r".\.venv-web\Scripts\python.exe scripts/evaluate_image_models.py"),
    ], widths=[1.7, 4.9])

    add_heading(doc, "8. Multimodal Search Method", 1)
    add_heading(doc, "8.1. Text Search", 2)
    add_body(doc, (
        "Text input is converted to lowercase, Unicode accents are removed, and configured Vietnamese phrases are mapped to English catalog terms. "
        "Longer phrases such as giày chạy bộ and xanh lá are processed before individual tokens. The retrieval score is the number of distinct normalized query tokens found in the product name, category, color, or description. "
        "A product is retained when at least one token matches."
    ))

    add_heading(doc, "8.2. Voice Search", 2)
    add_body(doc, (
        "The browser uses SpeechRecognition or webkitSpeechRecognition with vi-VN or en-US. It displays interim text, allows the customer to stop recording, and starts a search when final recognition completes. "
        "Only the transcript reaches the Python API; the server does not receive or store microphone audio. The CLI accepts a supplied transcript to reproduce the simulated voice requirement without depending on a microphone or online recognition service."
    ))

    add_heading(doc, "8.3. Image Similarity Search", 2)
    add_body(doc, (
        "The image pipeline center-crops RGB input to 224 by 224 pixels and applies the mean and standard deviation required by CLIP. "
        "The quantized CLIP ViT-B/32 vision encoder produces a 512-dimensional semantic vector. The vector is normalized to unit length. Catalog images pass through the same preprocessing and inference path when services start."
    ))
    add_body(doc, (
        "For a query vector q and product vector p, cosine similarity is score(q,p) = (q · p) / (||q|| ||p||). Because both vectors are normalized, the score is their dot product. "
        "The search compares the query with all 40 indexed products, then RankingService sorts the candidates and applies top_k. CLIP replaced MobileNetV2 class logits because semantic embeddings grouped the product categories more consistently."
    ))

    add_heading(doc, "8.4. Result Presentation", 2)
    add_body(doc, (
        "Every response preserves the original input and reports its mode, processing method, result count, backend duration, product data, rank, and score. "
        "Text and voice scores count matched tokens, while image scores are cosine similarities; the two scales therefore must not be compared directly."
    ))

    add_heading(doc, "9. Experimental Results", 1)
    add_heading(doc, "9.1. Test Set and Overall Results", 2)
    add_body(doc, (
        "The fixed evaluation set contains 12 labeled queries: five text queries, four voice transcripts, and three image queries. A query succeeds when its top-ranked product belongs to the declared relevant-product set. "
        "Eleven queries succeeded, producing an overall top-1 success rate of 91.67 percent."
    ))
    add_table(doc, ["Mode", "Queries", "Successful", "Top-1 success rate"], [
        ("Text", "5", "4", "80.00%"),
        ("Voice", "4", "4", "100.00%"),
        ("Image", "3", "3", "100.00%"),
        ("Overall", "12", "11", "91.67%"),
    ], widths=[1.5, 1.2, 1.4, 2.5])
    add_caption(doc, "Table 12. Top-1 retrieval results on the fixed evaluation set")
    add_body(doc, (
        "The only failed query was sneakers. The text method matches whole normalized tokens and does not currently expand sneakers to shoes, so it returned no candidate. "
        "This failure identifies a concrete vocabulary limitation rather than an unstable ranking result."
    ))

    add_heading(doc, "9.2. Image Model Comparison", 2)
    add_body(doc, (
        "A leave-one-out test measured whether the nearest different catalog photograph belonged to the same category. The exact query image was excluded from candidates. "
        "MobileNetV2 class logits achieved 29 of 40 correct categories. CLIP semantic embeddings achieved 38 of 40, an improvement of 22.5 percentage points. "
        "The stricter category-and-color result improved from 11 of 40 to 19 of 40."
    ))
    add_table(doc, ["Image representation", "Correct category", "Category accuracy", "Correct category and color"], [
        ("MobileNetV2 ImageNet logits", "29/40", "72.5%", "11/40 (27.5%)"),
        ("CLIP ViT-B/32 embedding", "38/40", "95.0%", "19/40 (47.5%)"),
    ], widths=[2.3, 1.35, 1.35, 1.6])
    add_caption(doc, "Table 13. Leave-one-out image-similarity comparison")
    add_body(doc, (
        "The category result shows that CLIP recognizes the product type more reliably. Color remains more difficult because backgrounds, lighting, clothing, and surrounding objects influence the visual embedding. "
        "The leave-one-out set is an internal regression proxy over 40 photographs and is not a claim of accuracy on unrestricted user images."
    ))

    add_heading(doc, "9.3. Demonstration Queries", 2)
    add_table(doc, ["Mode and input", "Processing", "Top result", "Score"], [
        ("Text: black shoes", "Bilingual normalization and whole-token matching", "Metro Black Shoes", "2.0000"),
        ("Voice: find running shoes", "Transcript validation, normalization, and text matching", "Sprint Blue Running Shoes", "2.0000"),
        ("Image: query_a.png", "CLIP 512-D embedding and cosine similarity", "Metro Black Shoes", "1.0000"),
    ], widths=[1.55, 2.65, 1.75, 0.65])
    add_caption(doc, "Table 14. Required text, voice, and image demonstrations")

    add_picture(doc, ASSETS / "text-search.png", 6.15)
    add_caption(doc, "Figure 4. Web demonstration of the text query black shoes with ranked products and scores")

    image_table = doc.add_table(rows=1, cols=2)
    image_table.autofit = False
    for cell, path in zip(image_table.rows[0].cells, [ASSETS / "voice-search.png", ASSETS / "image-search.png"]):
        cell.width = Inches(3.25)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=40, start=40, bottom=40, end=40)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(3.05))
    set_table_borders(image_table, color="E6E6E6", size="4")
    add_caption(doc, "Figure 5. Voice search with microphone input and image search with CLIP-based ranking")

    add_heading(doc, "9.4. Verification", 2)
    add_body(doc, (
        "The final verification run completed 36 automated tests successfully. The suite covers repository consistency, ranking, Vietnamese-English equivalence, Unicode normalization, microphone transcript flow, upload validation, CLIP embedding dimensions, catalog-image uniqueness, API boundaries, and leave-one-out image quality. "
        "A headless browser check also verified the 40-product catalog, text and voice results, image upload, diagnostic scores, product details, error states, 390-pixel mobile layout, and the absence of JavaScript errors."
    ))

    add_heading(doc, "10. Discussion", 1)
    add_heading(doc, "10.1. Strengths", 2)
    add_body(doc, (
        "The prototype maintains traceability from requirements and UML components to Python modules. All search modes produce the same query schema and ranked-product response, while their input-specific processing remains isolated behind services. "
        "The repository is balanced across four categories, photographs are unique and credited, and the web interface exposes the input, processing, returned products, and ranking score required by the assignment."
    ))
    add_body(doc, (
        "The implementation also completes several optional advanced features: microphone speech-to-text, real image embeddings, arbitrary image upload, bilingual search, and a responsive web interface. "
        "Replacing MobileNetV2 logits with CLIP changed only the image encoder and index construction, which demonstrates the extensibility intended by the three-layer design."
    ))

    add_heading(doc, "10.2. Limitations", 2)
    for item in [
        "Text retrieval uses a finite Vietnamese-English dictionary and whole-token matching; it does not provide general translation, stemming, or semantic synonym expansion.",
        "Browser microphone recognition depends on browser support, permission, a secure localhost or HTTPS context, and potentially an external recognition service.",
        "CLIP has not been fine-tuned for this catalog. It recognizes product categories well, but color ranking remains sensitive to background, lighting, and the portion of the image occupied by the product.",
        "The evaluation contains only 12 functional queries and 40 leave-one-out image cases. It is suitable for regression testing but too small for a broad accuracy claim.",
        "Prices have no declared currency, order-related use cases remain design-only, and the prototype does not persist accounts, sessions, carts, or transactions.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10.3. Recommended Improvements", 2)
    add_body(doc, (
        "The next image-search improvement should separate product type from color. CLIP can first retrieve products of the correct category, after which foreground segmentation and a Lab or HSV color descriptor can re-rank candidates by product color. "
        "Each product should also have several photographs with different angles and backgrounds. Text search would benefit from synonym expansion for terms such as sneakers and a small semantic-text model. "
        "A larger external query set should measure category, color, and exact-product retrieval separately."
    ))

    add_heading(doc, "11. Conclusion", 1)
    add_body(doc, (
        "The project delivers a working multimodal e-commerce search prototype that follows the required three-layer architecture. ProductRepository manages a 40-item catalog; QueryService creates consistent queries; SearchService performs text or image retrieval; RankingService produces deterministic top-k results; and the CLI and web adapters demonstrate the complete workflow."
    ))
    add_body(doc, (
        "The required text, simulated voice, image-similarity, and ranking functions are implemented and demonstrated with explicit input, processing, returned products, and scores. "
        "The web version extends the baseline with actual microphone capture, arbitrary image upload, Vietnamese-English text normalization, real CLIP embeddings, responsive presentation, and diagnostics. "
        "Automated tests and browser checks confirm the implementation, while the reported failures and image-color results define the current limits without overstating performance."
    ))

    add_heading(doc, "11.1. References", 2)
    for item in [
        "OpenAI. CLIP: Connecting text and images. https://github.com/openai/CLIP",
        "Xenova. CLIP ViT-B/32 ONNX model for Transformers.js. https://huggingface.co/Xenova/clip-vit-base-patch32",
        "ONNX Model Zoo. MobileNet image classification models. https://github.com/onnx/models/tree/main/validated/vision/classification/mobilenet",
        "MDN Web Docs. SpeechRecognition API. https://developer.mozilla.org/docs/Web/API/SpeechRecognition",
    ]:
        add_body(doc, item)


def normalize_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 6)
        style.paragraph_format.space_after = Pt(4)

    for paragraph in doc.paragraphs:
        if paragraph.style.name in {"Heading 1", "Heading 2"}:
            for run in paragraph.runs:
                set_font(run, size=14 if paragraph.style.name == "Heading 1" else 11.5, bold=True)

    for name in ("TOC 1", "TOC 2"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(8.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0

    for table_index, table in enumerate(doc.tables):
        if table_index == 0:
            continue
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        set_font(run, size=9)

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    update_existing_content(doc)
    append_sections(doc)
    normalize_styles(doc)
    doc.core_properties.title = "Multimodal E-Commerce Search System"
    doc.core_properties.subject = "Information System Analysis and Design Assignment 06"
    doc.core_properties.author = "Nguyen Tien Dat"
    doc.core_properties.comments = "Completed report aligned with the final Python implementation."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
