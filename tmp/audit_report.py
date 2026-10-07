from docx import Document

path = r"D:\My workspace\Project\HTTT-assignment06\output\docx\A6_V1_03_NguyenTienDat017.docx"
doc = Document(path)
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "images", len(doc.inline_shapes))
print("HEADINGS:")
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith("Heading"):
        print(paragraph.style.name, "|", paragraph.text)
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
print("STALE:")
for phrase in ["artificial feature", "predefined vector", "at least ten products", "simulated speech and image"]:
    print(phrase, text.lower().count(phrase.lower()))
